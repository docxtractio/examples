#!/usr/bin/env python3
# DocXtract Folder Watcher (Python)
# Date: 2026-09-01 | Update: 2026-09-07 | Author: Alok | File: folder_watcher.py
# Purpose: Iterate a local folder of invoices, extract each via the DocXtract API
#          (v3.1 defaults: document_type=invoice, model=default), write results to
#          Excel (one appending workbook or one per invoice), move processed files
#          to done/ and failures to failed/.
#
# Designed to be packaged as a single Windows .exe with PyInstaller:
#   pip install requests openpyxl pyinstaller
#   pyinstaller --onefile --name DocXtractFolderWatcher folder_watcher.py
#
# Configuration: config.ini next to the exe/script (created with defaults on first
# run), overridable per run with CLI flags. API key comes from config.ini or the
# DOCXTRACT_API_KEY environment variable.

import argparse
import configparser
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook, load_workbook

SUPPORTED = {".pdf", ".jpg", ".jpeg", ".png", ".webp"}

DEFAULT_CONFIG = {
    "api_url": "https://api.docxtract.io/v3.1/documents",
    "api_key": "",
    "in_folder": "invoices",
    "done_folder": "done",
    "failed_folder": "failed",
    "out_folder": "output",
    "excel_mode": "single",  # single | separate
    "model": "",          # blank = API default (prompt template "default")
    "document_type": "",  # blank = API default ("invoice")
}


def app_dir() -> Path:
    """Folder the exe/script lives in — config and relative folders resolve here."""
    if getattr(sys, "frozen", False):  # PyInstaller exe
        return Path(sys.executable).parent
    return Path(__file__).parent


def config_path() -> Path:
    return app_dir() / "config.ini"


def save_config(cfg: dict):
    parser = configparser.ConfigParser()
    parser["docxtract"] = {k: cfg.get(k, "") for k in DEFAULT_CONFIG}
    with open(config_path(), "w") as f:
        parser.write(f)
    if os.name == "posix":  # macOS/Linux: restrict config to the current user
        os.chmod(config_path(), 0o600)


def load_config() -> dict:
    path = config_path()
    parser = configparser.ConfigParser()
    cfg = dict(DEFAULT_CONFIG)
    if path.exists():
        parser.read(path)
        if parser.has_section("docxtract"):
            cfg.update({k: v for k, v in parser["docxtract"].items() if v})
    else:  # first run: write a template the user can edit
        save_config(cfg)
        print(f"Created {path}")
    return cfg


def ensure_api_key(cfg: dict, args: argparse.Namespace) -> str:
    """Resolve the API key: flag -> config.ini -> env var -> interactive prompt.
    A key entered at the prompt is saved to config.ini for future runs."""
    key = args.api_key or cfg.get("api_key") or os.environ.get("DOCXTRACT_API_KEY", "")
    if key:
        return key
    if not sys.stdin.isatty():  # scheduled/headless run: fail fast, never hang
        sys.exit("No API key. Set api_key in config.ini, DOCXTRACT_API_KEY, or pass --api-key.")
    key = input("Enter your DocXtract API key: ").strip()
    if not key:
        sys.exit("No API key entered.")
    cfg["api_key"] = key
    save_config(cfg)
    print(f"API key saved to {config_path()}")
    return key


def parse_args(cfg: dict) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="DocXtract folder-to-Excel invoice extraction")
    p.add_argument("--in-folder", default=cfg["in_folder"])
    p.add_argument("--done-folder", default=cfg["done_folder"])
    p.add_argument("--failed-folder", default=cfg["failed_folder"])
    p.add_argument("--out-folder", default=cfg["out_folder"])
    p.add_argument("--excel-mode", choices=["single", "separate"], default=cfg["excel_mode"])
    p.add_argument("--api-url", default=cfg["api_url"])
    p.add_argument("--api-key", default="")
    p.add_argument("--model", default=cfg["model"],
                   help="Prompt template name; blank = API default")
    p.add_argument("--document-type", default=cfg["document_type"],
                   help="Document type; blank = API default (invoice)")
    return p.parse_args()


def flatten(obj, prefix=""):
    """Flatten dynamic extraction JSON. Nested dicts -> dot-notation columns;
    lists of dicts (line_items etc.) -> separate item rows."""
    row, items = {}, []
    for key, val in obj.items():
        col = f"{prefix}.{key}" if prefix else key
        if isinstance(val, dict):
            r, i = flatten(val, col)
            row.update(r)
            items.extend(i)
        elif isinstance(val, list):
            if val and isinstance(val[0], dict):
                for el in val:
                    r, i = flatten(el)
                    items.append({"source_array": col, **r})
                    items.extend(i)
            else:
                row[col] = "; ".join(str(v) for v in val)
        else:
            row[col] = val
    return row, items


def extract(path: Path, api_url: str, api_key: str, model: str = "", document_type: str = "") -> dict:
    # Only send options that are actually set; when none are, the API applies
    # its v3.1 defaults (document_type=invoice, model=default).
    options = {}
    if model:
        options["model"] = model
    if document_type:
        options["document_type"] = document_type
    form_data = {"options": json.dumps(options)} if options else None
    with open(path, "rb") as f:
        resp = requests.post(
            api_url,
            headers={"Authorization": f"Bearer {api_key}"},
            files={"file": (path.name, f)},
            data=form_data,
            timeout=180,
        )
    result = resp.json()
    if not result.get("success"):
        err = result.get("error") or {}
        raise RuntimeError(err.get("message") or f"HTTP {resp.status_code}")
    return result["data"]


def move_unique(src: Path, target_dir: Path):
    dest = target_dir / src.name
    if dest.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = target_dir / f"{src.stem}_{stamp}{src.suffix}"
    src.rename(dest)


def read_sheet(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = list(rows[0])
    return [dict(zip(headers, r)) for r in rows[1:]]


def write_sheet(ws, rows: list[dict]):
    headers = []
    for r in rows:
        for k in r:
            if k not in headers:
                headers.append(k)
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])
    ws.freeze_panes = "A2"


def write_workbook(path: Path, invoice_rows: list[dict], item_rows: list[dict]):
    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb.create_sheet("Invoices"), invoice_rows)
    if item_rows:
        write_sheet(wb.create_sheet("LineItems"), item_rows)
    wb.save(path)


def main():
    cfg = load_config()
    args = parse_args(cfg)
    api_key = ensure_api_key(cfg, args)

    base = app_dir()
    folders = {
        name: (base / getattr(args, name)) if not Path(getattr(args, name)).is_absolute()
        else Path(getattr(args, name))
        for name in ("in_folder", "done_folder", "failed_folder", "out_folder")
    }
    for d in folders.values():
        d.mkdir(parents=True, exist_ok=True)

    files = sorted(p for p in folders["in_folder"].iterdir()
                   if p.is_file() and p.suffix.lower() in SUPPORTED)
    if not files:
        print(f"No PDF/image files found in {folders['in_folder']}. Nothing to do.")
        return

    print(f"Processing {len(files)} file(s) from {folders['in_folder']} (mode: {args.excel_mode})...")
    all_invoices, all_items = [], []
    ok = failed = 0

    for path in files:
        print(f"  {path.name} ... ", end="", flush=True)
        try:
            start = time.time()
            data = extract(path, args.api_url, api_key, args.model, args.document_type)
            row, items = flatten(data)
            invoice_row = {
                "source_file": path.name,
                "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                **row,
            }
            item_rows = [{"source_file": path.name, **i} for i in items]

            if args.excel_mode == "separate":
                write_workbook(folders["out_folder"] / f"{path.stem}.xlsx",
                               [invoice_row], item_rows)
            else:
                all_invoices.append(invoice_row)
                all_items.extend(item_rows)

            move_unique(path, folders["done_folder"])
            ok += 1
            print(f"OK ({time.time() - start:.1f}s)")
        except Exception as e:
            failed += 1
            print(f"FAILED — {e}")
            try:
                move_unique(path, folders["failed_folder"])
            except OSError:
                pass

    if args.excel_mode == "single" and all_invoices:
        xlsx = folders["out_folder"] / "invoices.xlsx"
        if xlsx.exists():  # merge with previous runs so the workbook keeps growing
            wb = load_workbook(xlsx, read_only=True)
            existing = read_sheet(wb["Invoices"]) if "Invoices" in wb.sheetnames else []
            existing_items = read_sheet(wb["LineItems"]) if "LineItems" in wb.sheetnames else []
            wb.close()
            all_invoices = existing + all_invoices
            all_items = existing_items + all_items
        write_workbook(xlsx, all_invoices, all_items)
        print(f"Workbook updated: {xlsx}")

    print(f"Done. {ok} succeeded, {failed} failed.")


if __name__ == "__main__":
    main()
